"""PPTX -> blocks. Slide number is the page; the title placeholder sets the
section. Shape geometry converts cleanly to points, so slides get real bboxes.
"""

from __future__ import annotations

import io

from pptx import Presentation

from app.services.ingestion.base import (
    ParsedBlock,
    ParsedDocument,
    SectionStack,
    classify,
    normalize_text,
)

EMU_PER_POINT = 12700


def _pt(emu: int | None) -> float:
    return round((emu or 0) / EMU_PER_POINT, 2)


def parse(data: bytes) -> ParsedDocument:
    prs = Presentation(io.BytesIO(data))
    page_w, page_h = _pt(prs.slide_width), _pt(prs.slide_height)
    sections = SectionStack()
    blocks: list[ParsedBlock] = []
    order = 0

    for slide_no, slide in enumerate(prs.slides, start=1):
        # python-pptx builds a fresh proxy on every `.title` access, so identity
        # comparison is useless here -- match on the stable shape id instead.
        title = slide.shapes.title
        title_id = title.shape_id if title is not None else None

        # Title first so the section is set before the body shapes are emitted.
        shapes = sorted(
            (s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()),
            key=lambda s: (s.shape_id != title_id, s.top or 0, s.left or 0),
        )

        for shape in shapes:
            is_title = shape.shape_id == title_id
            for para in shape.text_frame.paragraphs:
                text = normalize_text("".join(run.text for run in para.runs) or para.text)
                if not text:
                    continue
                if is_title:
                    sections.push(f"Slide {slide_no}: {text}", 1)

                blocks.append(
                    ParsedBlock(
                        page_no=slide_no,
                        section_path=sections.path,
                        order_idx=order,
                        block_type=classify(text, is_title),
                        text=text,
                        bbox={
                            "x0": _pt(shape.left),
                            "y0": _pt(shape.top),
                            "x1": _pt(shape.left) + _pt(shape.width),
                            "y1": _pt(shape.top) + _pt(shape.height),
                            "page_width": page_w,
                            "page_height": page_h,
                        },
                    )
                )
                order += 1

    return ParsedDocument(blocks=blocks, page_count=len(prs.slides))
