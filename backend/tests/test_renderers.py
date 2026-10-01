"""Renderer and format-spec tests.

Renderers are pure functions of ContentIR, so they can be tested exactly:
render, reopen with the library that owns the format, assert the structure
survived.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from pptx import Presentation

from app.schemas.content_ir import ContentIR, Node
from app.services.generation.formats import FORMATS, get_format
from app.services.rendering import (
    RENDERERS,
    UnknownRenderer,
    get_renderer,
    render_bytes,
    render_text,
)

ALL_KINDS = ContentIR(
    title="Incident CB/IR/2026/0412",
    nodes=[
        Node(
            id="n0",
            kind="callout",
            title="High severity",
            text="Credentials exposed.",
            severity="high",
            fact_ids=["fa"],
        ),
        Node(id="n1", kind="heading", text="Impact", level=2),
        Node(id="n2", kind="paragraph", text="37 credentials were compromised.", fact_ids=["fa"]),
        Node(id="n3", kind="bullets", items=["Enforce MFA", "Extend retention"], fact_ids=["fb"]),
        Node(
            id="n4", kind="table", rows=[["Metric", "Value"], ["Accounts", "37"]], fact_ids=["fa"]
        ),
        Node(id="n5", kind="quote", text="Contained within 6 hours.", fact_ids=["fb"]),
        Node(
            id="n6",
            kind="slide",
            title="Overview",
            items=["37 credentials"],
            notes="Speak slowly",
            fact_ids=["fa"],
        ),
        Node(id="n7", kind="post", items=["37 credentials compromised."], fact_ids=["fa"]),
        Node(
            id="n8",
            kind="panel",
            title="Credentials exposed",
            items=["37 credentials compromised"],
            notes="Single big number",
            fact_ids=["fa"],
        ),
        Node(
            id="n9",
            kind="scene",
            title="Opening",
            text="On 11 March 2026, 37 staff credentials were compromised.",
            items=["37 credentials"],
            notes="Title card over office footage",
            fact_ids=["fa"],
        ),
    ],
)

CITATIONS = {"fa": "incident_report.pdf (p.3, Findings)", "fb": "incident_report.pdf (p.4)"}
FULL_KWARGS = dict(include_citations=True, citations=CITATIONS, source_name="incident_report.pdf")


# --- format specs ----------------------------------------------------------


def test_all_nine_output_types_are_registered():
    assert set(FORMATS) == {
        "advisory",
        "ppt",
        "summary",
        "email",
        "linkedin",
        "press_release",
        "report",
        "infographic",
        "video",
    }


@pytest.mark.parametrize("key", sorted(FORMATS))
def test_every_spec_is_coherent(key):
    spec = get_format(key)
    assert spec.allowed_kinds, f"{key} allows no node kinds"
    assert spec.prompt_fragment.strip()
    assert spec.max_nodes > 0
    assert spec.renderers


@pytest.mark.parametrize("key", sorted(FORMATS))
def test_every_declared_renderer_actually_exists(key):
    """A spec promising a renderer that is not registered is a 500 waiting to happen."""
    for renderer in get_format(key).renderers:
        assert renderer in RENDERERS, f"{key} declares unknown renderer {renderer!r}"


def test_decks_only_allow_slides_and_posts_only_allow_posts():
    assert get_format("ppt").allowed_kinds == ("slide",)
    assert get_format("linkedin").allowed_kinds == ("post",)


def test_linkedin_is_a_single_post():
    assert get_format("linkedin").max_nodes == 1


def test_legacy_social_key_still_resolves_for_old_outputs():
    assert get_format("social") is get_format("linkedin")


# --- registry --------------------------------------------------------------


@pytest.mark.parametrize("key", sorted(RENDERERS))
def test_every_renderer_accepts_the_full_kwarg_set(key):
    """Regression: the export endpoint passes citations and source_name to every
    renderer, and one of them used to reject them with a 500."""
    assert render_bytes(ALL_KINDS, key, **FULL_KWARGS)


@pytest.mark.parametrize("key", sorted(RENDERERS))
def test_every_renderer_works_without_optional_kwargs(key):
    assert render_bytes(ALL_KINDS, key)


def test_render_text_refuses_binary_formats():
    with pytest.raises(UnknownRenderer, match="binary"):
        render_text(ALL_KINDS, "pptx")


def test_unknown_renderer_is_rejected():
    with pytest.raises(UnknownRenderer):
        render_bytes(ALL_KINDS, "hologram")


def test_extensions_and_media_types_agree_with_the_registry():
    assert get_renderer("docx").extension == "docx"
    assert "presentationml" in get_renderer("pptx").media_type
    assert get_renderer("markdown").media_type == "text/markdown"


@pytest.mark.parametrize("key", ["markdown", "text", "html"])
def test_text_renderers_are_byte_for_byte_reproducible(key):
    assert render_bytes(ALL_KINDS, key, **FULL_KWARGS) == render_bytes(
        ALL_KINDS, key, **FULL_KWARGS
    )


# --- pptx ------------------------------------------------------------------


def _deck(ir: ContentIR, **kw) -> Presentation:
    return Presentation(io.BytesIO(render_bytes(ir, "pptx", **kw)))


def test_pptx_is_widescreen():
    prs = _deck(ALL_KINDS)
    assert round(prs.slide_width.inches, 2) == 13.33
    assert round(prs.slide_height.inches, 2) == 7.5


def test_pptx_opens_with_a_title_slide_then_content():
    ir = ContentIR(
        title="Deck Title",
        nodes=[
            Node(id="n1", kind="slide", title="First", items=["a"], fact_ids=["fa"]),
            Node(id="n2", kind="slide", title="Second", items=["b"], fact_ids=["fa"]),
        ],
    )
    prs = _deck(ir, source_name="incident_report.pdf")
    assert len(prs.slides) == 3  # title + 2

    first = " ".join(s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame)
    assert "Deck Title" in first
    assert "incident_report.pdf" in first


def test_pptx_preserves_bullets_and_speaker_notes():
    ir = ContentIR(
        title="D",
        nodes=[
            Node(
                id="n1",
                kind="slide",
                title="Overview",
                items=["37 credentials compromised", "Contained in 6 hours"],
                notes="Mention the timeline",
                fact_ids=["fa"],
            ),
        ],
    )
    prs = _deck(ir)
    body = " ".join(s.text_frame.text for s in prs.slides[1].shapes if s.has_text_frame)
    assert "37 credentials compromised" in body
    assert "Contained in 6 hours" in body
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Mention the timeline"


def test_pptx_shows_a_source_footer_on_content_slides():
    prs = _deck(ALL_KINDS, source_name="incident_report.pdf")
    footer = " ".join(s.text_frame.text for s in prs.slides[1].shapes if s.has_text_frame)
    assert "incident_report.pdf" in footer


def test_pptx_folds_non_slide_nodes_into_slides():
    """Any ContentIR can become a deck, not just one built from slide nodes."""
    ir = ContentIR(
        title="D",
        nodes=[
            Node(id="n1", kind="heading", text="Findings", level=1),
            Node(id="n2", kind="bullets", items=["37 credentials", "4 offices"], fact_ids=["fa"]),
        ],
    )
    prs = _deck(ir)
    text = " ".join(
        s.text_frame.text for slide in prs.slides for s in slide.shapes if s.has_text_frame
    )
    assert "Findings" in text and "37 credentials" in text


def test_pptx_caps_bullets_per_slide():
    """A slide is a speaker prompt, not a document."""
    ir = ContentIR(
        title="D",
        nodes=[
            Node(
                id="n1",
                kind="slide",
                title="Many",
                items=[f"point {i}" for i in range(20)],
                fact_ids=["fa"],
            ),
        ],
    )
    body = " ".join(s.text_frame.text for s in _deck(ir).slides[1].shapes if s.has_text_frame)
    assert "point 19" not in body


def test_pptx_survives_an_empty_document():
    prs = _deck(ContentIR(title="Empty", nodes=[]))
    assert len(prs.slides) == 1


# --- docx ------------------------------------------------------------------


def _doc(ir: ContentIR, **kw) -> Document:
    return Document(io.BytesIO(render_bytes(ir, "docx", **kw)))


def test_docx_uses_real_word_styles():
    doc = _doc(ALL_KINDS)
    styles = {p.style.name for p in doc.paragraphs if p.text.strip()}
    assert "Title" in styles
    assert "List Bullet" in styles
    assert any(s.startswith("Heading") for s in styles)


def test_docx_renders_tables_as_tables():
    doc = _doc(ALL_KINDS)
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(0, 0).text == "Metric"
    assert doc.tables[0].cell(1, 1).text == "37"


def test_docx_citation_appendix_is_opt_in():
    without = [p.text.strip() for p in _doc(ALL_KINDS).paragraphs]
    assert "Sources" not in without

    with_sources = _doc(ALL_KINDS, include_citations=True, citations=CITATIONS)
    texts = [p.text.strip() for p in with_sources.paragraphs]
    assert "Sources" in texts
    assert any("incident_report.pdf (p.3, Findings)" in t for t in texts)


def test_docx_appendix_lists_each_fact_once():
    doc = _doc(ALL_KINDS, include_citations=True, citations=CITATIONS)
    texts = [p.text for p in doc.paragraphs]
    assert sum(1 for t in texts if t == CITATIONS["fa"]) == 1


def test_docx_carries_the_body_text():
    doc = _doc(ALL_KINDS)
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "37 credentials were compromised." in body
    assert "Enforce MFA" in body


# --- html ------------------------------------------------------------------


def test_html_is_a_standalone_document():
    out = render_text(ALL_KINDS, "html")
    assert out.startswith("<!doctype html>")
    assert "<title>Incident CB/IR/2026/0412</title>" in out
    assert "</html>" in out


def test_html_escapes_content():
    """Generated text is data, not markup."""
    ir = ContentIR(
        title="A & B",
        nodes=[
            Node(id="n1", kind="paragraph", text="<script>alert(1)</script>", fact_ids=["fa"]),
        ],
    )
    out = render_text(ir, "html")
    assert "<script>alert(1)</script>" not in out
    assert "&lt;script&gt;" in out
    assert "A &amp; B" in out


def test_html_renders_tables_and_callouts():
    out = render_text(ALL_KINDS, "html")
    assert "<table>" in out and "<td>Metric</td>" in out
    assert 'class="callout high"' in out


# --- markdown / text -------------------------------------------------------


def test_markdown_sources_appendix_uses_the_citation_map():
    out = render_text(ALL_KINDS, "markdown", include_citations=True, citations=CITATIONS)
    assert "## Sources" in out
    assert "- incident_report.pdf (p.3, Findings)" in out


def test_markdown_without_citations_has_no_appendix():
    assert "## Sources" not in render_text(ALL_KINDS, "markdown")


def test_text_renderer_stays_plain():
    out = render_text(ALL_KINDS, "text")
    assert "<" not in out
    assert "* Enforce MFA" in out
