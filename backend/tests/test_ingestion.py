"""Parser contract tests.

Traceability is the product, so these assert the anchors — page, section,
ordering, geometry — not just that text came out.
"""

from __future__ import annotations

import io
import pathlib

import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches, Pt

from app.services.ingestion import UnsupportedFormat, mime_for, parse_document
from app.services.ingestion.base import SectionStack, normalize_text

SAMPLE = pathlib.Path(__file__).resolve().parents[2] / "samples" / "incident_report.pdf"


# --- shared base -----------------------------------------------------------


def test_section_stack_nests_and_truncates():
    s = SectionStack()
    s.push("Findings", 1)
    s.push("Credential Exposure", 2)
    assert s.path == "Findings > Credential Exposure"
    # A new level-1 heading drops the deeper entries.
    s.push("Recommendations", 1)
    assert s.path == "Recommendations"


def test_normalize_text_collapses_runs():
    assert normalize_text("  a  b   c ") == "a b c"


def test_mime_rejects_unknown_extension():
    with pytest.raises(UnsupportedFormat):
        mime_for("notes.xlsx")


# --- PDF -------------------------------------------------------------------


@pytest.fixture(scope="module")
def sample_pdf() -> bytes:
    if not SAMPLE.exists():
        pytest.skip(f"{SAMPLE} missing; run scripts/make_samples.py")
    return SAMPLE.read_bytes()


def test_pdf_blocks_carry_full_provenance(sample_pdf: bytes):
    doc = parse_document("incident_report.pdf", sample_pdf)

    assert doc.page_count == 4
    assert doc.blocks, "expected blocks"

    for b in doc.blocks:
        assert 1 <= b.page_no <= doc.page_count
        assert b.text.strip()
        assert b.bbox is not None, "PDF blocks must have geometry"
        assert b.bbox["x0"] < b.bbox["x1"] and b.bbox["y0"] < b.bbox["y1"]
        assert 0 <= b.bbox["x0"] and b.bbox["x1"] <= b.bbox["page_width"] + 1

    # order_idx is dense and ascending: the UI relies on it for reading order.
    assert [b.order_idx for b in doc.blocks] == list(range(len(doc.blocks)))
    # Pages never go backwards.
    assert all(a.page_no <= b.page_no for a, b in zip(doc.blocks, doc.blocks[1:], strict=False))


def test_pdf_heading_hierarchy(sample_pdf: bytes):
    doc = parse_document("incident_report.pdf", sample_pdf)
    paths = {b.section_path for b in doc.blocks}
    assert "Cyber Security Incident Report > Findings > Credential Exposure" in paths
    assert "Cyber Security Incident Report > Recommendations" in paths
    # Nesting must actually nest, not flatten.
    assert any(p.count(" > ") == 2 for p in paths)


def test_pdf_bullets_classified_as_list_items(sample_pdf: bytes):
    doc = parse_document("incident_report.pdf", sample_pdf)
    bullets = [b for b in doc.blocks if b.block_type == "list_item"]
    assert len(bullets) >= 5
    # Base-14 Helvetica re-encodes U+2022 BULLET as U+00B7 MIDDLE DOT on
    # extraction, so accept either glyph.
    assert all(b.text.startswith(("\u2022", "\u00b7", "-", "*")) for b in bullets)


def test_pdf_key_claim_resolves_to_the_right_page_and_section(sample_pdf: bytes):
    """The claim the whole demo hangs on must land in Findings on page 3."""
    doc = parse_document("incident_report.pdf", sample_pdf)
    hits = [b for b in doc.blocks if b.text.startswith("37 user credentials were compromised")]
    assert len(hits) == 1
    block = hits[0]
    assert block.page_no == 3
    assert block.section_path.endswith("Findings > Credential Exposure")


def test_render_page_png(sample_pdf: bytes):
    from app.services.ingestion.pdf import render_page_png

    png = render_page_png(sample_pdf, 1, dpi=72)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    with pytest.raises(IndexError):
        render_page_png(sample_pdf, 99)


# --- DOCX ------------------------------------------------------------------


@pytest.fixture(scope="module")
def sample_docx() -> bytes:
    d = DocxDocument()
    d.add_heading("Advisory", level=1)
    d.add_paragraph("Body paragraph one.")
    d.add_heading("Findings", level=2)
    d.add_paragraph("37 credentials were compromised.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Accounts"
    table.cell(1, 1).text = "37"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_uses_styles_for_sections(sample_docx: bytes):
    doc = parse_document("advisory.docx", sample_docx)
    by_text = {b.text: b for b in doc.blocks}
    assert by_text["37 credentials were compromised."].section_path == "Advisory > Findings"
    assert by_text["Advisory"].block_type == "heading"
    # Word has no fixed pagination, so there is no geometry to report.
    assert all(b.bbox is None and b.page_no == 1 for b in doc.blocks)


def test_docx_captures_tables_in_order(sample_docx: bytes):
    doc = parse_document("advisory.docx", sample_docx)
    tables = [b for b in doc.blocks if b.block_type == "table"]
    assert len(tables) == 1
    assert "Metric | Value" in tables[0].text
    # The table comes after the paragraph that precedes it.
    assert tables[0].order_idx == max(b.order_idx for b in doc.blocks)


# --- PPTX ------------------------------------------------------------------


@pytest.fixture(scope="module")
def sample_pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Incident Overview"
    body = slide.placeholders[1].text_frame
    body.text = "37 credentials compromised"
    body.add_paragraph().text = "Contained within 6 hours"

    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "Recommendations"
    box = slide2.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = "Enforce MFA"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(18)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_pptx_slide_is_the_page_and_title_is_the_section(sample_pptx: bytes):
    doc = parse_document("deck.pptx", sample_pptx)
    assert doc.page_count == 2
    assert {b.page_no for b in doc.blocks} == {1, 2}

    body = next(b for b in doc.blocks if b.text == "37 credentials compromised")
    assert body.page_no == 1
    assert body.section_path == "Slide 1: Incident Overview"
    # Shape geometry converts to points and stays inside the slide.
    assert body.bbox is not None
    assert body.bbox["x1"] <= body.bbox["page_width"] + 1

    assert next(b for b in doc.blocks if b.text == "Enforce MFA").page_no == 2


def test_pptx_title_precedes_body_in_reading_order(sample_pptx: bytes):
    doc = parse_document("deck.pptx", sample_pptx)
    slide1 = [b for b in doc.blocks if b.page_no == 1]
    assert slide1[0].text == "Incident Overview"
    assert slide1[0].block_type == "heading"
