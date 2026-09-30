"""Generate the demo source documents in samples/.

A realistic incident report is worth more than lorem ipsum: it exercises
heading levels, bullets and — importantly for later phases — carries specific
numbers that the consistency checker will police.
"""

from __future__ import annotations

import pathlib

import pymupdf

OUT = pathlib.Path(__file__).resolve().parent.parent / "samples"

BODY, H3, H2, H1 = 10.5, 12.0, 14.0, 20.0
MARGIN_X, TOP, BOTTOM, LEADING = 62, 78, 760, 15.5

PAGES: list[list[tuple[float, str]]] = [
    [
        (H1, "Cyber Security Incident Report"),
        (BODY, "Reference: CB/IR/2026/0412 | Classification: Restricted"),
        (BODY, "Issued by the Computer Security Incident Response Team"),
        (H2, "Executive Summary"),
        (
            BODY,
            "On 11 March 2026 the Incident Response Team confirmed unauthorised access to the "
            "departmental identity provider. The intrusion was detected by anomalous "
            "authentication telemetry and contained within 6 hours of first alert.",
        ),
        (
            BODY,
            "A total of 37 user credentials were compromised across 4 regional offices. No "
            "evidence of data exfiltration beyond directory metadata has been identified at "
            "the time of writing.",
        ),
        (
            BODY,
            "The estimated financial impact is 18.4 lakh rupees, comprising incident response "
            "effort, forced credential rotation and third-party forensic support.",
        ),
        (H2, "Scope"),
        (
            BODY,
            "This report covers systems within the departmental Active Directory forest and "
            "the federated single sign-on gateway. Field office networks operating on isolated "
            "infrastructure are out of scope.",
        ),
    ],
    [
        (H2, "Incident Timeline"),
        (BODY, "All timestamps are Indian Standard Time (UTC+05:30)."),
        (BODY, "• 11 March 2026, 02:14 - Anomalous authentication volume detected."),
        (BODY, "• 11 March 2026, 02:51 - Security Operations Centre raises severity to High."),
        (BODY, "• 11 March 2026, 04:20 - Compromised service account disabled."),
        (BODY, "• 11 March 2026, 08:10 - Containment confirmed across all 4 offices."),
        (BODY, "• 12 March 2026, 11:00 - Forced password reset completed for 37 accounts."),
        (BODY, "• 14 March 2026, 17:30 - Forensic imaging of 2 domain controllers completed."),
        (H2, "Detection"),
        (
            BODY,
            "Detection originated from the authentication analytics rule AUTH-114, which flags "
            "credential stuffing patterns. Mean time to detect was 41 minutes from the first "
            "malicious authentication attempt.",
        ),
    ],
    [
        (H2, "Findings"),
        (H3, "Credential Exposure"),
        (
            BODY,
            "37 user credentials were compromised. Of these, 5 belonged to accounts holding "
            "elevated privileges within the identity management console. All 37 accounts were "
            "subject to mandatory password rotation.",
        ),
        (
            BODY,
            "Password reuse across the departmental portal and an external vendor service was "
            "the primary contributing factor in 29 of the 37 cases.",
        ),
        (H3, "Multi-Factor Authentication Gaps"),
        (
            BODY,
            "Multi-factor authentication was enforced on only 62 percent of in-scope accounts. "
            "Every one of the 37 compromised accounts lacked a second factor.",
        ),
        (H3, "Logging and Retention"),
        (
            BODY,
            "Authentication logs were retained for 30 days against a policy requirement of 180 "
            "days, limiting retrospective analysis of the intrusion.",
        ),
    ],
    [
        (H2, "Recommendations"),
        (
            BODY,
            "• Enforce multi-factor authentication on 100 percent of privileged accounts "
            "within 30 days.",
        ),
        (BODY, "• Extend authentication log retention to 180 days to meet policy."),
        (BODY, "• Deploy credential stuffing detection at the federation gateway."),
        (BODY, "• Complete a password hygiene audit across all 4 regional offices."),
        (BODY, "• Review vendor access agreements covering shared identity material."),
        (H2, "Conclusion"),
        (
            BODY,
            "The incident was contained without confirmed data loss, but the absence of "
            "multi-factor authentication on affected accounts represents a systemic control "
            "weakness. Implementation of the recommendations above is assessed as high priority.",
        ),
        (BODY, "Approved by the Chief Information Security Officer on 20 March 2026."),
    ],
]


def _wrap(text: str, size: float, width: float) -> list[str]:
    """Greedy wrap using real glyph widths so nothing overflows the margin."""
    font = pymupdf.Font("helv")
    words, lines, current = text.split(), [], ""
    for word in words:
        trial = f"{current} {word}".strip()
        if font.text_length(trial, size) <= width and current:
            current = trial
        elif not current:
            current = word
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def build_incident_report(path: pathlib.Path) -> None:
    doc = pymupdf.open()
    usable = 595 - 2 * MARGIN_X  # A4 width in points

    for page_no, entries in enumerate(PAGES, start=1):
        page = doc.new_page(width=595, height=842)
        y = TOP
        for size, text in entries:
            bold = size > BODY
            font = "hebo" if bold else "helv"
            for line in _wrap(text, size, usable):
                if y > BOTTOM:
                    break
                page.insert_text((MARGIN_X, y), line, fontsize=size, fontname=font)
                y += size * 1.35 if bold else LEADING
            y += 9 if bold else 6

        page.insert_text(
            (MARGIN_X, 800),
            f"CB/IR/2026/0412 — Page {page_no} of {len(PAGES)}",
            fontsize=8,
            fontname="helv",
            color=(0.45, 0.45, 0.45),
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()


def build_advisory_docx(path: pathlib.Path) -> None:
    """A short DOCX so the non-PDF ingestion path has a real fixture."""
    from docx import Document as Docx

    d = Docx()
    d.add_heading("Security Advisory CB/ADV/2026/07", level=1)
    d.add_paragraph("Issued 20 March 2026 | Classification: Restricted")
    d.add_heading("Summary", level=2)
    d.add_paragraph(
        "Following incident CB/IR/2026/0412, 37 user credentials were confirmed "
        "compromised. All affected accounts have been reset."
    )
    d.add_heading("Required Action", level=2)
    d.add_paragraph("Enforce multi-factor authentication on all privileged accounts.")
    d.add_paragraph("Extend authentication log retention to 180 days.")
    table = d.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Credentials compromised"
    table.cell(1, 1).text = "37"
    table.cell(2, 0).text = "Offices affected"
    table.cell(2, 1).text = "4"
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))


def build_briefing_pptx(path: pathlib.Path) -> None:
    """A short PPTX fixture covering the slide-as-page ingestion path."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Incident CB/IR/2026/0412"
    body = s1.placeholders[1].text_frame
    body.text = "37 credentials compromised across 4 offices"
    body.add_paragraph().text = "Contained within 6 hours of first alert"
    body.add_paragraph().text = "No confirmed data exfiltration"

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Recommendations"
    box = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3)).text_frame
    box.text = "Enforce MFA on 100 percent of privileged accounts"
    box.add_paragraph().text = "Extend log retention to 180 days"

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


if __name__ == "__main__":
    build_incident_report(OUT / "incident_report.pdf")
    build_advisory_docx(OUT / "security_advisory.docx")
    build_briefing_pptx(OUT / "incident_briefing.pptx")
    for f in sorted(OUT.iterdir()):
        print(f"wrote {f} ({f.stat().st_size:,} bytes)")
